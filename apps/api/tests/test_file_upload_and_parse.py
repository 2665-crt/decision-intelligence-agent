from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from universal_agent.main import app


def make_xlsx() -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["date", "region", "revenue"])
    worksheet.append(["2026-01-01", "east", 100])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def make_docx() -> BytesIO:
    document = Document()
    document.add_paragraph("项目风险说明")
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def test_excel_summary_and_explicit_selection() -> None:
    client = TestClient(app)
    task_id = client.post("/tasks").json()["id"]
    uploaded = client.post(
        f"/tasks/{task_id}/files",
        files={"file": ("sales.xlsx", make_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert uploaded.status_code == 201
    payload = uploaded.json()
    assert payload["parse_status"] == "succeeded"
    assert payload["summary"]["sheets"][0]["columns"] == ["date", "region", "revenue"]
    selection = client.post(f"/tasks/{task_id}/selections", json={"file_ids": [payload["id"]]})
    assert selection.status_code == 201
    assert selection.json()["file_ids"] == [payload["id"]]


def test_sixth_file_is_rejected_without_removing_first_five() -> None:
    client = TestClient(app)
    task_id = client.post("/tasks").json()["id"]
    for index in range(5):
        response = client.post(
            f"/tasks/{task_id}/files",
            files={"file": (f"sales-{index}.xlsx", make_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 201
    rejected = client.post(
        f"/tasks/{task_id}/files",
        files={"file": ("sales-6.xlsx", make_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert rejected.status_code == 422


def test_docx_text_is_marked_as_document_statement() -> None:
    client = TestClient(app)
    task_id = client.post("/tasks").json()["id"]
    response = client.post(
        f"/tasks/{task_id}/files",
        files={"file": ("brief.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 201
    assert response.json()["summary"]["text_evidence"][0]["level"] == "document_statement"
